"""Render a CandidateProfile as a CV document (text PDF or DOCX), in English or French.

Used for the synthetic part of the evaluation set: the profile (the answer key) is
written first and the document is generated from it, so the reference is exact by
construction. Dates are stored normalized ("2021-03") but displayed the way people
write them ("Mar 2021", "mars 2021"), which also exercises date normalization.
"""

import html
from pathlib import Path

from recruitment_ai.extraction.schemas import CandidateProfile

LABELS = {
    "en": {
        "summary": "Summary", "experience": "Experience", "education": "Education",
        "skills": "Skills", "certifications": "Certifications", "projects": "Projects",
        "languages": "Languages", "present": "Present", "dates_unknown": "",
    },
    "fr": {
        "summary": "Profil", "experience": "Expérience professionnelle", "education": "Formation",
        "skills": "Compétences", "certifications": "Certifications", "projects": "Projets",
        "languages": "Langues", "present": "aujourd'hui", "dates_unknown": "",
    },
}
MONTHS = {
    "en": ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"],
    "fr": ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."],
}
LAYOUTS = ("classic", "sidebar", "compact")


def display_date(value: str | None, lang: str) -> str:
    """'2021-03' -> 'Mar 2021' / 'mars 2021'; 'present' -> 'Present' / "aujourd'hui"; '2016' -> '2016'."""
    if not value:
        return ""
    if value == "present":
        return LABELS[lang]["present"]
    if len(value) == 7:
        year, month = value.split("-")
        return f"{MONTHS[lang][int(month) - 1]} {year}"
    return value


def date_range(start: str | None, end: str | None, lang: str) -> str:
    start_s, end_s = display_date(start, lang), display_date(end, lang)
    return f"{start_s} – {end_s}" if start_s and end_s else start_s or end_s


# --------------------------------------------------------------------------- PDF (HTML -> PyMuPDF Story)

_CSS = {
    "classic": """
        * {font-family: sans-serif; font-size: 10pt; color: #222;}
        h1 {font-size: 20pt; margin: 0;}
        .contact {color: #555; margin-bottom: 8pt;}
        h2 {font-size: 12pt; color: #1d4e89; border-bottom: 1px solid #1d4e89; margin-top: 10pt;}
        .role {font-weight: bold; margin-top: 6pt;}
        .dates {color: #555;}
        li {margin-bottom: 2pt;}
    """,
    "sidebar": """
        * {font-family: sans-serif; font-size: 9.5pt; color: #222;}
        h1 {font-size: 18pt; margin: 0;}
        h2 {font-size: 11pt; color: #7b2d26; margin-top: 8pt;}
        .side {background-color: #f1ede6; padding: 6pt;}
        .role {font-weight: bold; margin-top: 5pt;}
        .dates {color: #666; font-style: italic;}
    """,
    "compact": """
        * {font-family: serif; font-size: 9pt; color: #111;}
        h1 {font-size: 15pt; margin: 0; text-align: center;}
        .contact {text-align: center; margin-bottom: 4pt;}
        h2 {font-size: 10pt; text-transform: uppercase; margin-top: 6pt; margin-bottom: 2pt;}
        .role {font-weight: bold;}
        p, li {margin: 0;}
    """,
}


def _e(text: str | None) -> str:
    return html.escape(text or "")


def _bullets(description: str | None) -> str:
    if not description:
        return ""
    lines = [line.strip().lstrip("-").strip() for line in description.splitlines() if line.strip()]
    return "<ul>" + "".join(f"<li>{_e(line)}</li>" for line in lines) + "</ul>"


def _main_sections(p: CandidateProfile, lang: str) -> str:
    labels = LABELS[lang]
    parts = []
    if p.summary:
        parts.append(f"<h2>{labels['summary']}</h2><p>{_e(p.summary)}</p>")
    if p.experience:
        parts.append(f"<h2>{labels['experience']}</h2>")
        for x in p.experience:
            head = " — ".join(_e(v) for v in (x.title, x.company) if v)
            place = f", {_e(x.location)}" if x.location else ""
            dates = date_range(x.start_date, x.end_date, lang)
            parts.append(f"<p class='role'>{head}{place}</p>")
            if dates:
                parts.append(f"<p class='dates'>{_e(dates)}</p>")
            parts.append(_bullets(x.description))
    if p.education:
        parts.append(f"<h2>{labels['education']}</h2>")
        for ed in p.education:
            degree = ", ".join(_e(v) for v in (ed.degree, ed.field_of_study) if v)
            line = " — ".join(v for v in (degree, _e(ed.institution)) if v)
            dates = date_range(ed.start_date, ed.end_date, lang)
            grade = f" · {_e(ed.grade)}" if ed.grade else ""
            parts.append(f"<p class='role'>{line}</p><p class='dates'>{_e(dates)}{grade}</p>")
    if p.projects:
        parts.append(f"<h2>{labels['projects']}</h2>")
        for pr in p.projects:
            tech = f" ({_e(', '.join(pr.technologies))})" if pr.technologies else ""
            parts.append(f"<p><b>{_e(pr.title)}</b>{tech}: {_e(pr.description)}</p>")
    return "".join(parts)


def _side_sections(p: CandidateProfile, lang: str) -> str:
    labels = LABELS[lang]
    parts = []
    if p.skills:
        parts.append(f"<h2>{labels['skills']}</h2><p>{_e(', '.join(s.name for s in p.skills))}</p>")
    if p.certifications:
        items = "".join(
            f"<li>{_e(c.name)}{' (' + _e(display_date(c.issue_date, lang)) + ')' if c.issue_date else ''}</li>"
            for c in p.certifications
        )
        parts.append(f"<h2>{labels['certifications']}</h2><ul>{items}</ul>")
    if p.languages:
        langs = ", ".join(
            f"{_e(x.language)} ({_e(x.proficiency)})" if x.proficiency else _e(x.language) for x in p.languages
        )
        parts.append(f"<h2>{labels['languages']}</h2><p>{langs}</p>")
    return "".join(parts)


def profile_to_html(p: CandidateProfile, lang: str, layout: str) -> str:
    contact = " · ".join(_e(v) for v in (p.location, p.email, p.phone) if v)
    header = f"<h1>{_e(p.full_name)}</h1><p class='contact'>{contact}</p>"
    main, side = _main_sections(p, lang), _side_sections(p, lang)
    if layout == "sidebar":
        return (
            f"{header}<table><tr>"
            f"<td class='side' width='32%' valign='top'>{side}</td>"
            f"<td width='68%' valign='top'>{main}</td></tr></table>"
        )
    return header + main + side


def render_pdf(p: CandidateProfile, path: Path, lang: str = "en", layout: str = "classic") -> None:
    """Write a text-based (not scanned) A4 PDF of the profile."""
    import pymupdf

    if layout not in LAYOUTS:
        raise ValueError(f"unknown layout {layout!r}, expected one of {LAYOUTS}")
    path.parent.mkdir(parents=True, exist_ok=True)
    story = pymupdf.Story(html=profile_to_html(p, lang, layout), user_css=_CSS[layout])
    writer = pymupdf.DocumentWriter(str(path))
    page = pymupdf.paper_rect("a4")
    area = page + (40, 40, -40, -40)
    more = True
    while more:
        device = writer.begin_page(page)
        more, _ = story.place(area)
        story.draw(device)
        writer.end_page()
    writer.close()


# --------------------------------------------------------------------------- DOCX


def render_docx(p: CandidateProfile, path: Path, lang: str = "en") -> None:
    """Write a Word document of the profile (single column, standard styles)."""
    from docx import Document

    labels = LABELS[lang]
    doc = Document()
    doc.add_heading(p.full_name or "", level=0)
    doc.add_paragraph(" · ".join(v for v in (p.location, p.email, p.phone) if v))
    if p.summary:
        doc.add_heading(labels["summary"], level=1)
        doc.add_paragraph(p.summary)
    if p.experience:
        doc.add_heading(labels["experience"], level=1)
        for x in p.experience:
            head = " — ".join(v for v in (x.title, x.company) if v) + (f", {x.location}" if x.location else "")
            run = doc.add_paragraph().add_run(head)
            run.bold = True
            dates = date_range(x.start_date, x.end_date, lang)
            if dates:
                doc.add_paragraph(dates)
            for line in (x.description or "").splitlines():
                if line.strip():
                    doc.add_paragraph(line.strip().lstrip("-").strip(), style="List Bullet")
    if p.education:
        doc.add_heading(labels["education"], level=1)
        for ed in p.education:
            degree = ", ".join(v for v in (ed.degree, ed.field_of_study) if v)
            doc.add_paragraph(" — ".join(v for v in (degree, ed.institution) if v)).runs[0].bold = True
            extra = date_range(ed.start_date, ed.end_date, lang) + (f" · {ed.grade}" if ed.grade else "")
            if extra.strip():
                doc.add_paragraph(extra)
    if p.projects:
        doc.add_heading(labels["projects"], level=1)
        for pr in p.projects:
            tech = f" ({', '.join(pr.technologies)})" if pr.technologies else ""
            doc.add_paragraph(f"{pr.title}{tech}: {pr.description or ''}")
    if p.skills:
        doc.add_heading(labels["skills"], level=1)
        doc.add_paragraph(", ".join(s.name for s in p.skills))
    if p.certifications:
        doc.add_heading(labels["certifications"], level=1)
        for c in p.certifications:
            date = f" ({display_date(c.issue_date, lang)})" if c.issue_date else ""
            doc.add_paragraph(f"{c.name}{date}", style="List Bullet")
    if p.languages:
        doc.add_heading(labels["languages"], level=1)
        doc.add_paragraph(", ".join(
            f"{x.language} ({x.proficiency})" if x.proficiency else x.language for x in p.languages
        ))
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
