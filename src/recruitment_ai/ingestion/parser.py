"""Document ingestion: CV files (text PDF, scanned PDF, image, DOCX, TXT) -> text.

Every parser returns NFKC-normalized text: PDFs draw "fi"/"fl" as ligature glyphs
("Certiﬁed") and some CVs use full-width or compatibility characters; NFKC turns them back
into plain letters while keeping French accents intact.
"""

import unicodedata
from abc import ABC, abstractmethod
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field

from recruitment_ai.ingestion.ocr import OCREngine, default_engine, page_image
from recruitment_ai.utils.logging import get_logger

logger = get_logger(__name__)

# A page with fewer extractable characters than this is treated as a scan and OCR'd
MIN_TEXT_CHARS_PER_PAGE = 30


def normalize_text(text: str) -> str:
    return unicodedata.normalize("NFKC", text or "")


class ParsedPage(BaseModel):
    """Extracted content from an individual document page."""

    page_number: int
    text: str
    layout_info: dict[str, Any] = Field(default_factory=dict)


class ParsedDocument(BaseModel):
    """Complete parsed document output."""

    filename: str
    file_type: str
    total_pages: int
    raw_text: str
    pages: list[ParsedPage] = Field(default_factory=list)
    metadata: dict[str, Any] = Field(default_factory=dict)


class DocumentParser(ABC):
    """Abstract interface for file parsers."""

    @abstractmethod
    def parse(self, file_path: Path) -> ParsedDocument:
        """Parse given document file and return structured text."""


class PdfParser(DocumentParser):
    """PDF parser: PyMuPDF text layer, with OCR for pages that have none (scans)."""

    def __init__(self, ocr: OCREngine | None = None) -> None:
        self._ocr = ocr

    @property
    def ocr(self) -> OCREngine:
        return self._ocr or default_engine()

    def parse(self, file_path: Path) -> ParsedDocument:
        import pymupdf

        path = Path(file_path)
        pages: list[ParsedPage] = []
        with pymupdf.open(str(path)) as doc:
            for index, page in enumerate(doc):
                text = page.get_text("text") or ""
                method, confidence = "text_layer", None
                if len(text.strip()) < MIN_TEXT_CHARS_PER_PAGE:
                    result = self.ocr.image_to_text(page_image(page))
                    text, method, confidence = result.text, "ocr", round(result.mean_confidence, 3)
                pages.append(ParsedPage(
                    page_number=index + 1, text=normalize_text(text),
                    layout_info={"rect": list(page.rect), "method": method, "ocr_confidence": confidence},
                ))
        methods = {p.layout_info["method"] for p in pages}
        return ParsedDocument(
            filename=path.name,
            file_type="pdf",
            total_pages=len(pages),
            raw_text="\n\n".join(p.text for p in pages),
            pages=pages,
            metadata={
                "engine": "PyMuPDF+OCR" if "ocr" in methods else "PyMuPDF",
                "input_type": "scanned_pdf" if methods == {"ocr"} else "text_pdf" if methods == {"text_layer"} else "mixed_pdf",
            },
        )


class DocxParser(DocumentParser):
    """DOCX parser using python-docx (paragraphs, then table rows)."""

    def parse(self, file_path: Path) -> ParsedDocument:
        import docx

        path = Path(file_path)
        doc = docx.Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        full_text = normalize_text("\n".join(paragraphs))
        return ParsedDocument(
            filename=path.name,
            file_type="docx",
            total_pages=1,
            raw_text=full_text,
            pages=[ParsedPage(page_number=1, text=full_text)],
            metadata={"engine": "python-docx", "input_type": "docx"},
        )


class ImageOCRParser(DocumentParser):
    """Image file (photo or scan of a CV) parser with OCR."""

    def __init__(self, ocr: OCREngine | None = None) -> None:
        self._ocr = ocr

    def parse(self, file_path: Path) -> ParsedDocument:
        import numpy as np
        from PIL import Image

        path = Path(file_path)
        image = np.asarray(Image.open(path).convert("RGB"))
        result = (self._ocr or default_engine()).image_to_text(image)
        text = normalize_text(result.text)
        return ParsedDocument(
            filename=path.name,
            file_type=path.suffix.lstrip(".").lower(),
            total_pages=1,
            raw_text=text,
            pages=[ParsedPage(page_number=1, text=text,
                              layout_info={"method": "ocr", "ocr_confidence": round(result.mean_confidence, 3)})],
            metadata={"engine": "RapidOCR PP-OCRv5", "input_type": "image"},
        )


class UnifiedDocumentParser(DocumentParser):
    """Router parser that chooses the appropriate parser from the file extension."""

    def __init__(self, ocr: OCREngine | None = None) -> None:
        self.pdf_parser = PdfParser(ocr)
        self.docx_parser = DocxParser()
        self.ocr_parser = ImageOCRParser(ocr)

    def parse(self, file_path: Path) -> ParsedDocument:
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == ".pdf":
            return self.pdf_parser.parse(path)
        if ext == ".docx":
            return self.docx_parser.parse(path)
        if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
            return self.ocr_parser.parse(path)
        if ext in {".txt", ".md"}:
            text = normalize_text(path.read_text(encoding="utf-8", errors="ignore"))
            return ParsedDocument(
                filename=path.name,
                file_type=ext.lstrip("."),
                total_pages=1,
                raw_text=text,
                pages=[ParsedPage(page_number=1, text=text)],
                metadata={"engine": "plain_text", "input_type": "text"},
            )
        raise ValueError(f"Unsupported document format: {ext}")
